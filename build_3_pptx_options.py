import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

LOGO_PATH = r"c:\FTH\FTH-D365_FO\fujitsu_wordmark_clean.png"
IMG_S1 = r"c:\FTH\FTH-D365_FO\assets\copilot_d365_hero_fluent.jpg"
IMG_S2 = r"c:\FTH\FTH-D365_FO\assets\paradigm_shift_3d.jpg"
IMG_S3 = r"c:\FTH\FTH-D365_FO\assets\fujitsu_copilot_value_3d.jpg"
IMG_S4 = r"c:\FTH\FTH-D365_FO\assets\live_demo_portal_3d.jpg"

SCRIPTS = {
    1: (
        "【Slide 1: Executive Welcome & Strategic Purpose】\n"
        "Tone: Formal, Confident, Inspiring | Target Duration: ~40 seconds\n\n"
        "\"Good morning, members of the executive leadership team, esteemed colleagues, and distinguished guests.\n\n"
        "On behalf of Fujitsu Thailand, it is my distinct privilege to welcome you to today’s executive keynote: "
        "'Empowering Enterprise ERP with Microsoft Copilot AI.'\n\n"
        "As organizations navigate accelerating digital transformation and tightening tax regulations in Thailand, "
        "traditional ERP systems must evolve beyond passive data entry. Today, we present Fujitsu’s strategic vision "
        "for Dynamics 365 Finance & Operations—natively enhanced by Microsoft Copilot Studio and Azure AI.\n\n"
        "Our objective today is clear: to demonstrate how autonomous AI streamlines Accounts Payable and Thai tax compliance, "
        "while creating proprietary, high-margin Intellectual Property that solidifies Fujitsu Thailand as the premier "
        "digital transformation partner in the region.\n\n"
        "Thank you for your time, and we look forward to an insightful session.\""
    ),
    2: (
        "【Slide 2: Industry Shift & Local Regulatory Challenges】\n"
        "Tone: Strategic, Analytical, Problem-Solving | Target Duration: ~45 seconds\n\n"
        "\"Turning to Slide 2, let us examine the operational realities and regulatory drivers shaping corporate finance today.\n\n"
        "Historically, ERP systems functioned as Systems of Record—reliable at logging historical transactions, but fundamentally passive. "
        "In practice, finance and Accounts Payable teams across Thailand continue to bear an immense manual burden: spending over 70% "
        "of their time scrutinizing hundreds of PDF tax invoices, matching purchase orders, and validating 13-digit Tax IDs line by line.\n\n"
        "Compounding this challenge is Thailand’s rapidly tightening compliance environment. The Revenue Department’s ongoing transition "
        "toward mandatory e-Tax Invoices, coupled with intricate Withholding Tax classifications, means that even minor clerical typos "
        "risk delayed filings and statutory penalties.\n\n"
        "This friction represents a transformative opportunity. By elevating Dynamics 365 into a System of Intelligence, "
        "we deploy proactive AI agents that pre-validate data, match transactions, and flag discrepancies before posting—converting "
        "finance from a transactional cost center into a strategic growth driver.\""
    ),
    3: (
        "【Slide 3: Product Excellence & Fujitsu Strategic Value】\n"
        "Tone: Visionary, High-Impact, Balanced | Target Duration: ~50 seconds\n\n"
        "\"On Slide 3, we illustrate how this initiative generates balanced value across both the Dynamics 365 product platform "
        "and Fujitsu Thailand’s commercial business.\n\n"
        "From a product perspective, we natively embed Microsoft Copilot Studio and Azure AI into Dynamics 365 via Dataverse, "
        "guaranteeing zero external data leakage. Copilot functions as an autonomous drafting co-pilot: reading invoices, checking 3-way matching, "
        "and pre-populating draft journals—slashing cycle times by over 70%. Crucially, we enforce strict Human-in-the-Loop governance: "
        "no transaction is committed to the ledger without explicit approval from an authorized accountant.\n\n"
        "From Fujitsu Thailand’s strategic perspective, this establishes a decisive first-mover advantage. By productizing our localization "
        "expertise into reusable, AI-infused Thai Localization IP, we accelerate implementation velocity while capturing higher consulting margins. "
        "This offering addresses the acute compliance needs of Japanese multinationals and major domestic enterprises, reinforcing Fujitsu’s legacy "
        "of Japanese engineering quality, robust governance, and trusted delivery.\""
    ),
    4: (
        "【Slide 4: Executive Agenda & Transition to Live Solution Pitch】\n"
        "Tone: Engaging, Action-Oriented, Inviting | Target Duration: ~30 seconds\n\n"
        "\"Finally, on Slide 4, let us outline today’s interactive agenda.\n\n"
        "Following this opening briefing, we will transition directly into our production-grade Live Interactive Pitch Platform, "
        "which walks through four core pillars across twelve executive slides:\n"
        "1. First, the official Microsoft Copilot Studio reference architecture, detailing enterprise security, Dataverse integration, and multi-agent coordination;\n"
        "2. Second, an in-depth walkthrough of our top five Thai localization use cases—spanning AP Invoice Intelligence, Thai Tax Copilot, e-Tax Gateway pre-validation, Continuous Reconciliation, and Policy Q&A;\n"
        "3. Third, a pragmatic 8-week Proof of Concept delivery roadmap, engineered to deliver measurable ROI and verified business value within two months; and\n"
        "4. Fourth, an interactive presentation experience equipped with dual-host narration and real-time scenario simulation.\n\n"
        "Without further delay, I invite our executive team to transition with me into the live interactive platform to witness the future of Dynamics 365 and Copilot AI in action.\n\n"
        "Please join me as we launch the live demonstration.\""
    )
}

# Common Colors
BG_WHITE = RGBColor(255, 255, 255)
BG_SLATE_50 = RGBColor(248, 250, 252)
CARD_BORDER = RGBColor(226, 232, 240)
ACCENT_RED = RGBColor(220, 38, 38)
ACCENT_BLUE = RGBColor(2, 132, 199)
ACCENT_PURPLE = RGBColor(124, 58, 237)
ACCENT_EMERALD = RGBColor(5, 150, 105)
TEXT_DARK = RGBColor(15, 23, 42)
TEXT_SUB = RGBColor(71, 85, 105)
TEXT_MUTED = RGBColor(148, 163, 184)

# =============================================================================
# OPTION 1: MICROSOFT COPILOT FLUENT KEYNOTE (Split-Screen 3D Visual + Keywords)
# =============================================================================
def build_option_1():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_base(badge, title, slide_num):
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_WHITE
        bg.line.fill.background()

        top_r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
        top_r.fill.solid()
        top_r.fill.fore_color.rgb = ACCENT_RED
        top_r.line.fill.background()

        if os.path.exists(LOGO_PATH):
            slide.shapes.add_picture(LOGO_PATH, Inches(11.4), Inches(0.35), width=Inches(1.35))

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.3), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = badge.upper()
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_RED
        p0.space_after = Pt(2)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK

        ftr = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
        ftf = ftr.text_frame
        ftf.margin_left = ftf.margin_top = 0
        p = ftf.paragraphs[0]
        p.text = f"OPTION 1: COPILOT FLUENT KEYNOTE • SLIDE {slide_num} OF 4"
        p.font.name = "Segoe UI"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED

        return slide

    # Slide 1 (Option 1)
    s1 = add_base("Fujitsu Thailand (FTH) • Executive Keynote", "Empowering Enterprise ERP with Microsoft Copilot AI", 1)
    if os.path.exists(IMG_S1):
        s1.shapes.add_picture(IMG_S1, Inches(0.8), Inches(1.7), width=Inches(6.4))
    
    # Right Side: Big Keyword Badges
    kw_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(1.7), Inches(5.033), Inches(5.05))
    kw_box.fill.solid()
    kw_box.fill.fore_color.rgb = BG_SLATE_50
    kw_box.line.color.rgb = CARD_BORDER
    kw_box.line.width = Pt(1.5)
    ktf = kw_box.text_frame
    ktf.word_wrap = True
    ktf.margin_left = ktf.margin_right = Inches(0.4)
    ktf.margin_top = Inches(0.35)

    kp0 = ktf.paragraphs[0]
    kp0.text = "🎯 CORE STRATEGIC PILLARS"
    kp0.font.name = "Segoe UI"
    kp0.font.size = Pt(12)
    kp0.font.bold = True
    kp0.font.color.rgb = ACCENT_BLUE
    kp0.space_after = Pt(14)

    s1_cards = [
        ("⚡ AUTONOMOUS FINANCE", "System of Intelligence • >70% AP Cycle Reduction", ACCENT_BLUE),
        ("🏛️ 100% THAI COMPLIANT", "VAT 7% • WHT 1-5% • 50 Tawi • e-Tax Ready", ACCENT_PURPLE),
        ("🏆 STRATEGIC FTH IP ASSET", "Proprietary Repeatable Solution • Japanese MNC Trust", ACCENT_RED)
    ]
    for h, b, col in s1_cards:
        p_h = ktf.add_paragraph()
        p_h.text = h
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = col
        p_h.space_after = Pt(2)

        p_b = ktf.add_paragraph()
        p_b.text = b
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(10.5)
        p_b.font.color.rgb = TEXT_SUB
        p_b.space_after = Pt(14)

    s1.notes_slide.notes_text_frame.text = SCRIPTS[1]

    # Slide 2 (Option 1)
    s2 = add_base("Market Drivers & Local Complexity", "The Shift: From System of Record to System of Intelligence", 2)
    if os.path.exists(IMG_S2):
        s2.shapes.add_picture(IMG_S2, Inches(0.8), Inches(1.7), width=Inches(6.4))

    kw_box2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(1.7), Inches(5.033), Inches(5.05))
    kw_box2.fill.solid()
    kw_box2.fill.fore_color.rgb = BG_SLATE_50
    kw_box2.line.color.rgb = CARD_BORDER
    kw_box2.line.width = Pt(1.5)
    ktf2 = kw_box2.text_frame
    ktf2.word_wrap = True
    ktf2.margin_left = ktf2.margin_right = Inches(0.4)
    ktf2.margin_top = Inches(0.35)

    kp0 = ktf2.paragraphs[0]
    kp0.text = "📊 OPERATIONAL REALITIES & DRIVERS"
    kp0.font.name = "Segoe UI"
    kp0.font.size = Pt(12)
    kp0.font.bold = True
    kp0.font.color.rgb = ACCENT_RED
    kp0.space_after = Pt(14)

    s2_cards = [
        ("⚠️ 70% MANUAL OVERLOAD", "Repetitive invoice verification • High month-end burnout", ACCENT_RED),
        ("🏛️ MANDATORY e-TAX COMPLIANCE", "Revenue Dept XML schemas • Zero tolerance for typos", ACCENT_PURPLE),
        ("🚀 PROACTIVE INTELLIGENCE", "Pre-posting validation • Continuous automated subledger audit", ACCENT_BLUE)
    ]
    for h, b, col in s2_cards:
        p_h = ktf2.add_paragraph()
        p_h.text = h
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = col
        p_h.space_after = Pt(2)

        p_b = ktf2.add_paragraph()
        p_b.text = b
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(10.5)
        p_b.font.color.rgb = TEXT_SUB
        p_b.space_after = Pt(14)

    s2.notes_slide.notes_text_frame.text = SCRIPTS[2]

    # Slide 3 (Option 1)
    s3 = add_base("Product Excellence & Strategic Value", "Strengthening Dynamics 365 & Maximizing Fujitsu Value", 3)
    if os.path.exists(IMG_S3):
        s3.shapes.add_picture(IMG_S3, Inches(0.8), Inches(1.7), width=Inches(6.4))

    kw_box3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(1.7), Inches(5.033), Inches(5.05))
    kw_box3.fill.solid()
    kw_box3.fill.fore_color.rgb = BG_SLATE_50
    kw_box3.line.color.rgb = CARD_BORDER
    kw_box3.line.width = Pt(1.5)
    ktf3 = kw_box3.text_frame
    ktf3.word_wrap = True
    ktf3.margin_left = ktf3.margin_right = Inches(0.4)
    ktf3.margin_top = Inches(0.35)

    kp0 = ktf3.paragraphs[0]
    kp0.text = "💎 DUAL-ENGINE VALUE PROPOSITION"
    kp0.font.name = "Segoe UI"
    kp0.font.size = Pt(12)
    kp0.font.bold = True
    kp0.font.color.rgb = ACCENT_PURPLE
    kp0.space_after = Pt(14)

    s3_cards = [
        ("🔒 ZERO DATA LEAKAGE", "Copilot Studio + Dataverse • 100% Secure Enterprise Tenant", ACCENT_BLUE),
        ("🛡️ HUMAN-IN-THE-LOOP", "AI drafts & pre-validates • Accountants retain full sign-off", ACCENT_EMERALD),
        ("🚀 HIGH-MARGIN REUSABLE IP", "First-mover localization moat • Repeatable commercial asset", ACCENT_RED)
    ]
    for h, b, col in s3_cards:
        p_h = ktf3.add_paragraph()
        p_h.text = h
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = col
        p_h.space_after = Pt(2)

        p_b = ktf3.add_paragraph()
        p_b.text = b
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(10.5)
        p_b.font.color.rgb = TEXT_SUB
        p_b.space_after = Pt(14)

    s3.notes_slide.notes_text_frame.text = SCRIPTS[3]

    # Slide 4 (Option 1)
    s4 = add_base("Solution Roadmap & Live Demonstration", "Executive Agenda & Transition to Live Solution Pitch", 4)
    if os.path.exists(IMG_S4):
        s4.shapes.add_picture(IMG_S4, Inches(0.8), Inches(1.7), width=Inches(6.4))

    kw_box4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(1.7), Inches(5.033), Inches(5.05))
    kw_box4.fill.solid()
    kw_box4.fill.fore_color.rgb = RGBColor(236, 253, 245) # Emerald tint
    kw_box4.line.color.rgb = ACCENT_EMERALD
    kw_box4.line.width = Pt(1.5)
    ktf4 = kw_box4.text_frame
    ktf4.word_wrap = True
    ktf4.margin_left = ktf4.margin_right = Inches(0.4)
    ktf4.margin_top = Inches(0.35)

    kp0 = ktf4.paragraphs[0]
    kp0.text = "🚀 READY FOR LIVE DEMONSTRATION"
    kp0.font.name = "Segoe UI"
    kp0.font.size = Pt(12)
    kp0.font.bold = True
    kp0.font.color.rgb = ACCENT_EMERALD
    kp0.space_after = Pt(12)

    s4_steps = [
        ("1. Official Architecture", "Copilot Studio + Dataverse Multi-Agent framework"),
        ("2. Top 5 Thai Use Cases", "AP Invoice, Tax Copilot, e-Tax Gateway & Recon"),
        ("3. 8-Week POC Delivery", "Fast ROI • Low-risk implementation roadmap"),
        ("4. Interactive Experience", "Scenario simulations & dual-host executive speech")
    ]
    for h, b in s4_steps:
        p_h = ktf4.add_paragraph()
        p_h.text = f"• {h}"
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(11.5)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_DARK

        p_b = ktf4.add_paragraph()
        p_b.text = f"   {b}"
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = TEXT_SUB
        p_b.space_after = Pt(8)

    p_cta = ktf4.add_paragraph()
    p_cta.text = "🌐 Direct Platform Link: https://wititgithub.github.io/fth-d365-fo-copilot-deck/"
    p_cta.font.name = "Segoe UI"
    p_cta.font.size = Pt(9.5)
    p_cta.font.bold = True
    p_cta.font.color.rgb = ACCENT_BLUE

    s4.notes_slide.notes_text_frame.text = SCRIPTS[4]

    out_p = r"c:\FTH\FTH-D365_FO\FTH_D365_Opening_Option1_Copilot_Fluent.pptx"
    prs.save(out_p)
    print(f"Saved Option 1 -> {out_p}")


# =============================================================================
# OPTION 2: MICROSOFT EXECUTIVE ARCHITECTURE & METRIC CARDS (Clean Grid)
# =============================================================================
def build_option_2():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_base(badge, title, slide_num):
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_WHITE
        bg.line.fill.background()

        top_r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
        top_r.fill.solid()
        top_r.fill.fore_color.rgb = ACCENT_BLUE
        top_r.line.fill.background()

        if os.path.exists(LOGO_PATH):
            slide.shapes.add_picture(LOGO_PATH, Inches(11.4), Inches(0.35), width=Inches(1.35))

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.3), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0

        p0 = tf.paragraphs[0]
        p0.text = badge.upper()
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_BLUE
        p0.space_after = Pt(2)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK

        ftr = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
        ftf = ftr.text_frame
        ftf.margin_left = ftf.margin_top = 0
        p = ftf.paragraphs[0]
        p.text = f"OPTION 2: EXECUTIVE MINIMALIST & METRICS • SLIDE {slide_num} OF 4"
        p.font.name = "Segoe UI"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED

        return slide

    # Helper: 4 Metric Cards Across Bottom
    def add_metric_cards(slide, cards_data):
        card_w = Inches(2.7)
        gap = Inches(0.31)
        for i, (metric, label, desc, col) in enumerate(cards_data):
            left_pos = Inches(0.8) + i * (card_w + gap)
            c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(4.3), card_w, Inches(2.45))
            c.fill.solid()
            c.fill.fore_color.rgb = BG_SLATE_50
            c.line.color.rgb = CARD_BORDER
            c.line.width = Pt(1.2)

            stripe = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(4.3), card_w, Inches(0.06))
            stripe.fill.solid()
            stripe.fill.fore_color.rgb = col
            stripe.line.fill.background()

            tf = c.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.2)
            tf.margin_top = Inches(0.18)

            p_m = tf.paragraphs[0]
            p_m.text = metric
            p_m.font.name = "Segoe UI"
            p_m.font.size = Pt(26)
            p_m.font.bold = True
            p_m.font.color.rgb = col
            p_m.space_after = Pt(2)

            p_l = tf.add_paragraph()
            p_l.text = label.upper()
            p_l.font.name = "Segoe UI"
            p_l.font.size = Pt(10)
            p_l.font.bold = True
            p_l.font.color.rgb = TEXT_DARK
            p_l.space_after = Pt(4)

            p_d = tf.add_paragraph()
            p_d.text = desc
            p_d.font.name = "Segoe UI"
            p_d.font.size = Pt(8.5)
            p_d.font.color.rgb = TEXT_SUB

    # Slide 1 (Option 2)
    s1 = add_base("Fujitsu Thailand (FTH) • Executive Briefing", "Empowering Enterprise ERP with Microsoft Copilot AI", 1)
    if os.path.exists(IMG_S1):
        s1.shapes.add_picture(IMG_S1, Inches(0.8), Inches(1.6), width=Inches(11.733), height=Inches(2.5))
    add_metric_cards(s1, [
        (">70%", "Cycle Reduction", "Autonomous AP extraction & 3-way matching", ACCENT_BLUE),
        ("100%", "Thai Compliance", "PP.30 VAT, PND 3/53, 50 Tawi & e-Tax XML", ACCENT_PURPLE),
        ("ZERO", "Data Leakage", "Grounded in secure Microsoft Dataverse", ACCENT_EMERALD),
        ("HIGH", "Strategic Margin", "Repeatable enterprise IP asset for FTH", ACCENT_RED)
    ])
    s1.notes_slide.notes_text_frame.text = SCRIPTS[1]

    # Slide 2 (Option 2)
    s2 = add_base("Market Challenges & Transformation", "The Shift: From System of Record to System of Intelligence", 2)
    if os.path.exists(IMG_S2):
        s2.shapes.add_picture(IMG_S2, Inches(0.8), Inches(1.6), width=Inches(11.733), height=Inches(2.5))
    add_metric_cards(s2, [
        ("70%+", "Manual Burden", "Unproductive time spent reviewing paper/PDFs", ACCENT_RED),
        ("13-DIGIT", "Tax ID Precision", "Automatic verification eliminates filing errors", ACCENT_PURPLE),
        ("e-TAX", "Mandate Ready", "Pre-submission XML validation gate", ACCENT_BLUE),
        ("PROACTIVE", "Continuous Audit", "Discrepancies caught before month-end close", ACCENT_EMERALD)
    ])
    s2.notes_slide.notes_text_frame.text = SCRIPTS[2]

    # Slide 3 (Option 2)
    s3 = add_base("Strategic Technology & Market Fusion", "Strengthening Dynamics 365 & Maximizing Fujitsu Value", 3)
    if os.path.exists(IMG_S3):
        s3.shapes.add_picture(IMG_S3, Inches(0.8), Inches(1.6), width=Inches(11.733), height=Inches(2.5))
    add_metric_cards(s3, [
        ("NATIVE", "Copilot Fusion", "Deep integration into standard D365 processes", ACCENT_BLUE),
        ("APPROVAL", "Human-in-the-Loop", "Accountants retain full journal sign-off authority", ACCENT_EMERALD),
        ("1ST MOVER", "Market Lead", "Premier AI Thai Localization packaged solution", ACCENT_RED),
        ("JAPANESE", "Quality Standard", "Fujitsu enterprise governance & delivery trust", ACCENT_PURPLE)
    ])
    s3.notes_slide.notes_text_frame.text = SCRIPTS[3]

    # Slide 4 (Option 2)
    s4 = add_base("Demonstration Roadmap & Handover", "Executive Agenda & Transition to Live Solution Pitch", 4)
    if os.path.exists(IMG_S4):
        s4.shapes.add_picture(IMG_S4, Inches(0.8), Inches(1.6), width=Inches(11.733), height=Inches(2.5))
    add_metric_cards(s4, [
        ("STAGE 1", "Copilot Arch", "Multi-Agent Dataverse framework (Slide 3)", ACCENT_BLUE),
        ("STAGE 2", "5 Use Cases", "AP, Thai Tax, e-Tax, Reconciliation (Slide 4)", ACCENT_PURPLE),
        ("8 WEEKS", "POC Delivery", "Low-risk, high-velocity ROI verification", ACCENT_EMERALD),
        ("LIVE NOW", "Platform Launch", "Interactive demonstration ready to begin", ACCENT_RED)
    ])
    s4.notes_slide.notes_text_frame.text = SCRIPTS[4]

    out_p = r"c:\FTH\FTH-D365_FO\FTH_D365_Opening_Option2_Executive_Minimalist.pptx"
    prs.save(out_p)
    print(f"Saved Option 2 -> {out_p}")


# =============================================================================
# OPTION 3: VISIONARY SPOTLIGHT (Centerpiece 3D Visual + Floating Pill Badges)
# =============================================================================
def build_option_3():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_base(badge, title, slide_num):
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_WHITE
        bg.line.fill.background()

        top_r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
        top_r.fill.solid()
        top_r.fill.fore_color.rgb = ACCENT_PURPLE
        top_r.line.fill.background()

        if os.path.exists(LOGO_PATH):
            slide.shapes.add_picture(LOGO_PATH, Inches(11.4), Inches(0.35), width=Inches(1.35))

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.3), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = 0

        p0 = tf.paragraphs[0]
        p0.text = badge.upper()
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9.5)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_PURPLE
        p0.space_after = Pt(2)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_DARK

        ftr = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
        ftf = ftr.text_frame
        ftf.margin_left = ftf.margin_top = 0
        p = ftf.paragraphs[0]
        p.text = f"OPTION 3: VISIONARY SPOTLIGHT • SLIDE {slide_num} OF 4"
        p.font.name = "Segoe UI"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED

        return slide

    def build_spotlight_slide(slide, img_path, left_pill, right_pill, bottom_cue):
        # Center Spotlight Visual
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, Inches(3.6), Inches(1.7), width=Inches(6.133))

        # Left Floating Pill Card
        c_left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1), Inches(2.6), Inches(3.8))
        c_left.fill.solid()
        c_left.fill.fore_color.rgb = BG_SLATE_50
        c_left.line.color.rgb = CARD_BORDER
        c_left.line.width = Pt(1.5)
        ltf = c_left.text_frame
        ltf.word_wrap = True
        ltf.margin_left = ltf.margin_right = Inches(0.2)
        ltf.margin_top = Inches(0.2)
        lp0 = ltf.paragraphs[0]
        lp0.text = left_pill[0]
        lp0.font.name = "Segoe UI"
        lp0.font.size = Pt(12)
        lp0.font.bold = True
        lp0.font.color.rgb = left_pill[1]
        lp0.space_after = Pt(8)
        for bullet in left_pill[2]:
            pb = ltf.add_paragraph()
            pb.text = f"• {bullet}"
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(9.5)
            pb.font.color.rgb = TEXT_SUB
            pb.space_after = Pt(6)

        # Right Floating Pill Card
        c_right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.933), Inches(2.1), Inches(2.6), Inches(3.8))
        c_right.fill.solid()
        c_right.fill.fore_color.rgb = BG_SLATE_50
        c_right.line.color.rgb = CARD_BORDER
        c_right.line.width = Pt(1.5)
        rtf = c_right.text_frame
        rtf.word_wrap = True
        rtf.margin_left = rtf.margin_right = Inches(0.2)
        rtf.margin_top = Inches(0.2)
        rp0 = rtf.paragraphs[0]
        rp0.text = right_pill[0]
        rp0.font.name = "Segoe UI"
        rp0.font.size = Pt(12)
        rp0.font.bold = True
        rp0.font.color.rgb = right_pill[1]
        rp0.space_after = Pt(8)
        for bullet in right_pill[2]:
            pb = rtf.add_paragraph()
            pb.text = f"• {bullet}"
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(9.5)
            pb.font.color.rgb = TEXT_SUB
            pb.space_after = Pt(6)

        # Bottom Key Cue Banner
        cue_b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.2), Inches(6.15), Inches(8.933), Inches(0.55))
        cue_b.fill.solid()
        cue_b.fill.fore_color.rgb = RGBColor(241, 245, 249)
        cue_b.line.color.rgb = CARD_BORDER
        cue_b.line.width = Pt(1)
        ctf = cue_b.text_frame
        ctf.word_wrap = True
        ctf.margin_top = Inches(0.12)
        p = ctf.paragraphs[0]
        p.text = bottom_cue
        p.font.name = "Segoe UI"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        p.alignment = PP_ALIGN.CENTER

    # Slide 1 (Option 3)
    s1 = add_base("Executive Keynote Opening", "Empowering Enterprise ERP with Microsoft Copilot AI", 1)
    build_spotlight_slide(
        s1, IMG_S1,
        ("⚡ AUTONOMOUS ERP", ACCENT_BLUE, ["Automates AP Review", "Reads Thai & English", "Creates Draft Journals"]),
        ("🏛️ THAI COMPLIANCE", ACCENT_PURPLE, ["PP.30 VAT 7%", "WHT 1-5% / 50 Tawi", "e-Tax XML Validation"]),
        "🎯 Strategic Objective: Delivering autonomous efficiency while building high-margin IP for Fujitsu Thailand."
    )
    s1.notes_slide.notes_text_frame.text = SCRIPTS[1]

    # Slide 2 (Option 3)
    s2 = add_base("Market Challenges & Opportunity", "From System of Record to System of Intelligence", 2)
    build_spotlight_slide(
        s2, IMG_S2,
        ("⚠️ PAIN POINTS", ACCENT_RED, [">70% Manual Workload", "13-Digit Tax ID Typos", "Month-End Stress"]),
        ("🚀 AI OPPORTUNITY", ACCENT_BLUE, ["Pre-Validation Gate", "Autonomous Matching", "Zero Audit Exposure"]),
        "💡 Key Insight: Transforming Accounts Payable from an operational bottleneck into a proactive financial command center."
    )
    s2.notes_slide.notes_text_frame.text = SCRIPTS[2]

    # Slide 3 (Option 3)
    s3 = add_base("Dual Value Creation", "Strengthening Dynamics 365 & Maximizing Fujitsu Value", 3)
    build_spotlight_slide(
        s3, IMG_S3,
        ("💎 D365 STRENGTH", ACCENT_BLUE, ["Copilot Studio Native", "Zero Data Leakage", "Human-in-the-Loop"]),
        ("🏆 FTH VALUE", ACCENT_RED, ["First-Mover in Thailand", "Reusable Commercial IP", "Japanese Quality Trust"]),
        "🎯 Mutual Impact: Enterprise clients gain 70% efficiency; Fujitsu captures scalable high-margin IP leadership."
    )
    s3.notes_slide.notes_text_frame.text = SCRIPTS[3]

    # Slide 4 (Option 3)
    s4 = add_base("Demonstration Handover", "Executive Agenda & Transition to Live Solution Pitch", 4)
    build_spotlight_slide(
        s4, IMG_S4,
        ("📑 12-SLIDE DECK", ACCENT_PURPLE, ["Reference Architecture", "Top 5 Thai Use Cases", "8-Week POC Delivery"]),
        ("🚀 LIVE DEMO", ACCENT_EMERALD, ["Interactive Scenarios", "Dual-Host Studio Audio", "Immediate Launch"]),
        "🌐 Live Demonstration Ready: Transitioning into the production-ready 12-slide interactive solution platform."
    )
    s4.notes_slide.notes_text_frame.text = SCRIPTS[4]

    out_p = r"c:\FTH\FTH-D365_FO\FTH_D365_Opening_Option3_Visionary_Spotlight.pptx"
    prs.save(out_p)
    print(f"Saved Option 3 -> {out_p}")


if __name__ == "__main__":
    build_option_1()
    build_option_2()
    build_option_3()
    # Also set Option 1 as default master FTH_D365_Executive_Opening_Deck.pptx
    import shutil
    shutil.copy(r"c:\FTH\FTH-D365_FO\FTH_D365_Opening_Option1_Copilot_Fluent.pptx", r"c:\FTH\FTH-D365_FO\FTH_D365_Executive_Opening_Deck.pptx")
    shutil.copy(r"c:\FTH\FTH-D365_FO\FTH_D365_Opening_Option1_Copilot_Fluent.pptx", r"c:\FTH\FTH-D365_FO\assets\FTH_D365_Executive_Opening_Deck.pptx")
    print("Master FTH_D365_Executive_Opening_Deck.pptx updated with Option 1 (Split-Screen High-Impact Visuals).")
