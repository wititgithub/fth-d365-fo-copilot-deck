import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Executive Light Corporate Palette (Fujitsu Executive Standard)
    BG_PAGE = RGBColor(255, 255, 255)         # Pure White Canvas
    CARD_BG = RGBColor(248, 250, 252)         # Crisp Slate-50 Card Fill
    CARD_BORDER = RGBColor(226, 232, 240)     # Clean Slate-200 Border
    
    # Brand Accents
    ACCENT_RED = RGBColor(220, 38, 38)        # Fujitsu Red (#DC2626)
    ACCENT_BLUE = RGBColor(2, 132, 199)       # Tech Azure (#0284C7)
    ACCENT_PURPLE = RGBColor(124, 58, 237)    # AI Indigo (#7C3AED)
    ACCENT_EMERALD = RGBColor(5, 150, 105)    # Success Emerald (#059669)
    ACCENT_NAVY = RGBColor(30, 58, 138)       # Enterprise Navy (#1E3A8A)

    # Typography Colors
    TEXT_TITLE = RGBColor(15, 23, 42)         # Deep Charcoal Slate-900 (High contrast)
    TEXT_SUBTITLE = RGBColor(71, 85, 105)     # Slate-600
    TEXT_BODY_HEAD = RGBColor(30, 41, 59)     # Slate-800
    TEXT_BODY = RGBColor(71, 85, 105)         # Slate-600
    TEXT_MUTED = RGBColor(148, 163, 184)      # Slate-400

    LOGO_PATH = r"c:\FTH\FTH-D365_FO\fujitsu_wordmark_clean.png"
    blank_layout = prs.slide_layouts[6]

    def add_slide_scaffolding(badge_text, title_text, subtitle_text, slide_num):
        slide = prs.slides.add_slide(blank_layout)

        # 1. Pure White Background
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_PAGE
        bg_shape.line.fill.background()

        # 2. Top Red Hairline Rule (Fujitsu Signature Accent)
        top_rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
        top_rule.fill.solid()
        top_rule.fill.fore_color.rgb = ACCENT_RED
        top_rule.line.fill.background()

        # 3. Fujitsu Logo at Top Right
        if os.path.exists(LOGO_PATH):
            slide.shapes.add_picture(LOGO_PATH, Inches(11.35), Inches(0.35), width=Inches(1.4))

        # 4. Header Text Block
        hdr_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(10.3), Inches(1.2))
        tf = hdr_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Category Badge
        p_badge = tf.paragraphs[0]
        p_badge.text = badge_text.upper()
        p_badge.font.name = "Segoe UI"
        p_badge.font.size = Pt(9.5)
        p_badge.font.bold = True
        p_badge.font.color.rgb = ACCENT_RED
        p_badge.space_after = Pt(2)

        # Action Title
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_TITLE
        p_title.space_after = Pt(2)

        # Subtitle / Strategic Takeaway
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.name = "Segoe UI"
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = TEXT_SUBTITLE

        # 5. Hairline Divider below Header
        div_rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.733), Inches(0.015))
        div_rule.fill.solid()
        div_rule.fill.fore_color.rgb = CARD_BORDER
        div_rule.line.fill.background()

        # 6. Professional Slide Footer
        ftr_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
        ftf = ftr_box.text_frame
        ftf.word_wrap = True
        ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
        p_ftr = ftf.paragraphs[0]
        p_ftr.text = f"FUJITSU THAILAND (FTH) • DYNAMICS 365 & COPILOT AI EXECUTIVE KEYNOTE • SLIDE {slide_num} OF 4"
        p_ftr.font.name = "Segoe UI"
        p_ftr.font.size = Pt(8.5)
        p_ftr.font.bold = True
        p_ftr.font.color.rgb = TEXT_MUTED

        return slide

    # =========================================================================
    # SLIDE 1: Title & Executive Welcome (English White Theme)
    # =========================================================================
    s1 = add_slide_scaffolding(
        "Fujitsu Thailand (FTH) • Strategic Executive Keynote",
        "Empowering Enterprise ERP with Microsoft Copilot AI",
        "Advancing Dynamics 365 Finance & Operations into Autonomous AI & Unlocking Strategic Value for Fujitsu Thailand",
        1
    )

    card_w = Inches(3.64)
    card_h = Inches(4.35)
    top_pos = Inches(1.85)
    gap = Inches(0.4)

    s1_pillars = [
        ("⚡ 1. Autonomous ERP", ACCENT_BLUE, [
            ("Workflow Automation", "Transforms manual accounting workflows into an intelligent system that reads, verifies, and pre-drafts transactions."),
            ("Native D365 Fusion", "Direct integration with Copilot Studio, Azure AI, and D365 F&O via Dataverse—zero external data leakage."),
            ("Operational Efficiency", "Slashes repetitive Accounts Payable cycle time by over 70%, freeing teams for strategic financial analysis.")
        ]),
        ("🏛️ 2. Thai Tax Compliance", ACCENT_PURPLE, [
            ("Full Regulatory Support", "Purpose-built for Revenue Department requirements: PP.30 VAT 7%, PND 3/53, and 50 Tawi certificates."),
            ("e-Tax Pre-Validation", "Automatic validation gate checking 13-digit Tax IDs and 5-digit branch codes prior to XML generation."),
            ("Continuous Subledger Audit", "Monitors General Ledger against AP, AR, and Tax subledgers throughout the month, eliminating closing panic.")
        ]),
        ("🏆 3. Strategic Value for FTH", ACCENT_RED, [
            ("Proprietary IP Leadership", "Establishes FTH as the premier partner with a proven, repeatable AI-infused Thai Localization package."),
            ("High-Margin Reusability", "Packaged solution scales across Japanese MNCs and major domestic corporate enterprises in Thailand."),
            ("Trusted Global Partner", "Combines Fujitsu’s renowned Japanese quality and governance with Microsoft’s premier enterprise AI platform.")
        ])
    ]

    for i, (title, color, items) in enumerate(s1_pillars):
        left_pos = Inches(0.8) + i * (card_w + gap)
        c = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        c.line.width = Pt(1.2)
        
        # Top colored accent stripe inside card
        stripe = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card_w, Inches(0.08))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = color
        stripe.line.fill.background()

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(12)

        for heading, body in items:
            pb = tf.add_paragraph()
            pb.text = f"• {heading}: "
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(10.5)
            pb.font.bold = True
            pb.font.color.rgb = TEXT_BODY_HEAD

            run = pb.add_run()
            run.text = body
            run.font.name = "Segoe UI"
            run.font.size = Pt(9.5)
            run.font.bold = False
            run.font.color.rgb = TEXT_BODY
            pb.space_after = Pt(8)

    # Slide 1 Bottom Presenter Banner
    b_bar = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.52))
    b_bar.fill.solid()
    b_bar.fill.fore_color.rgb = RGBColor(241, 245, 249)
    b_bar.line.color.rgb = CARD_BORDER
    b_bar.line.width = Pt(1)
    btf = b_bar.text_frame
    btf.word_wrap = True
    btf.margin_top = Inches(0.12)
    bp = btf.paragraphs[0]
    bp.text = "🎤 Presenter: Senior Solution Principal & Executive Host, Fujitsu Thailand (FTH)   |   ⏱️ Opening Duration: ~2 Minutes   |   ● Direct Live Solution Pitch Handover"
    bp.font.name = "Segoe UI"
    bp.font.size = Pt(10)
    bp.font.bold = True
    bp.font.color.rgb = TEXT_SUBTITLE
    bp.alignment = PP_ALIGN.CENTER

    # Slide 1 Formal English Executive Script (In Notes)
    s1.notes_slide.notes_text_frame.text = (
        "【Slide 1: Executive Welcome & Visionary Opening】\n"
        "Tone: Formal, Confident, Prestigious | Target Duration: ~40 seconds\n\n"
        "\"Good morning, members of the executive leadership team, esteemed colleagues, and distinguished guests.\n\n"
        "On behalf of Fujitsu Thailand, it is my distinct privilege to welcome you to today’s executive briefing: "
        "'Empowering Enterprise ERP with Microsoft Copilot AI.'\n\n"
        "As organizations navigate rapid digital evolution and intensifying regulatory demands in Thailand, traditional "
        "enterprise resource planning systems must evolve beyond passive data recording. Today, we are proud to present "
        "Fujitsu’s strategic vision for Dynamics 365 Finance & Operations—natively enhanced by Microsoft Copilot Studio and Azure AI.\n\n"
        "Our objective this morning is twofold: first, to demonstrate how autonomous AI streamlines everyday accounting, accounts payable, "
        "and Thai tax compliance; and second, to illustrate how this initiative creates proprietary, high-margin Intellectual Property "
        "that solidifies Fujitsu Thailand’s position as the premier digital transformation partner in the region.\n\n"
        "Thank you for your time, and we look forward to an insightful session.\""
    )

    # =========================================================================
    # SLIDE 2: Industry Shift - System of Record to System of Intelligence
    # =========================================================================
    s2 = add_slide_scaffolding(
        "Market Drivers & Business Challenges",
        "The Paradigm Shift: From System of Record to System of Intelligence",
        "Navigating Thai financial and tax complexities, and unlocking Fujitsu’s breakthrough opportunity.",
        2
    )

    s2_panels = [
        ("⚠️ Traditional ERP Limits", ACCENT_RED, [
            ("Manual Overload", "Accounts Payable teams spend >70% of time manually verifying PDF invoices, receipts, and line items."),
            ("Human Error Risks", "Typos in 13-digit Tax IDs or 5-digit branch codes trigger costly month-end tax filing discrepancies."),
            ("Month-End Bottlenecks", "Volume spikes create severe month-end stress, delaying PP.30 filings and obscuring cash visibility.")
        ]),
        ("🏛️ Strict Thai Mandates", ACCENT_PURPLE, [
            ("e-Tax Enforcement", "Revenue Department enforces strict XML / e-Filing schemas requiring 100% pre-validation before submission."),
            ("Complex WHT Rules", "Variable 1%, 2%, 3%, 5% rates require exact categorization based on the underlying nature of services."),
            ("Penalty Exposure", "Late or erroneous filings lead to surcharges, statutory interest, and unnecessary audit friction.")
        ]),
        ("🚀 Fujitsu's Opportunity", ACCENT_BLUE, [
            ("System of Intelligence", "Elevates D365 to analyze, pre-validate, and predict discrepancies before journal entries are posted."),
            ("Human-in-the-Loop", "AI prepares validated drafts; finance professionals retain exclusive approval authority."),
            ("Market Leadership", "FTH delivers modern AI innovation that traditional ERP consulting competitors in Thailand cannot match.")
        ])
    ]

    for i, (title, color, items) in enumerate(s2_panels):
        left_pos = Inches(0.8) + i * (card_w + gap)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card_w, card_h)
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = CARD_BORDER
        c.line.width = Pt(1.2)

        stripe = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, card_w, Inches(0.08))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = color
        stripe.line.fill.background()

        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(12)

        for heading, body in items:
            pb = tf.add_paragraph()
            pb.text = f"• {heading}: "
            pb.font.name = "Segoe UI"
            pb.font.size = Pt(10.5)
            pb.font.bold = True
            pb.font.color.rgb = TEXT_BODY_HEAD

            run = pb.add_run()
            run.text = body
            run.font.name = "Segoe UI"
            run.font.size = Pt(9.5)
            run.font.bold = False
            run.font.color.rgb = TEXT_BODY
            pb.space_after = Pt(8)

    # Slide 2 Bottom Banner
    b_bar2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.52))
    b_bar2.fill.solid()
    b_bar2.fill.fore_color.rgb = RGBColor(241, 245, 249)
    b_bar2.line.color.rgb = CARD_BORDER
    b_bar2.line.width = Pt(1)
    btf2 = b_bar2.text_frame
    btf2.word_wrap = True
    btf2.margin_top = Inches(0.12)
    bp2 = btf2.paragraphs[0]
    bp2.text = "💡 Executive Takeaway: Evolving ERP from passive data recording into proactive, AI-assisted financial governance."
    bp2.font.name = "Segoe UI"
    bp2.font.size = Pt(10)
    bp2.font.bold = True
    bp2.font.color.rgb = ACCENT_NAVY
    bp2.alignment = PP_ALIGN.CENTER

    s2.notes_slide.notes_text_frame.text = (
        "【Slide 2: Industry Shift & Local Regulatory Challenges】\n"
        "Tone: Problem-Solving, Strategic | Target Duration: ~45 seconds\n\n"
        "\"Turning to Slide 2, let us examine the operational realities and regulatory drivers shaping modern corporate finance.\n\n"
        "Historically, ERP systems were designed as Systems of Record—reliable at logging historical transactions, but fundamentally passive. "
        "In practice, our finance and Accounts Payable teams across Thailand continue to bear an immense manual burden. "
        "More than 70% of an AP accountant's time is consumed by repetitive administrative tasks: scrutinizing hundreds of PDF tax invoices, "
        "cross-referencing purchase orders, and validating 13-digit Tax Identification Numbers and branch codes line by line.\n\n"
        "Compounding this challenge is Thailand’s rapidly tightening regulatory landscape. The Revenue Department’s ongoing transition "
        "toward mandatory e-Tax Invoices and e-Receipts, coupled with intricate Withholding Tax classifications, means that even a minor "
        "typographical error can stall monthly PP.30 filings or invite statutory penalties.\n\n"
        "This operational friction represents a transformative opportunity for Fujitsu. By elevating Dynamics 365 into a System of Intelligence, "
        "we introduce proactive AI agents that validate data, match transactions, and flag discrepancies before posting—converting the finance "
        "department from a transactional cost center into an agile, strategic growth driver.\""
    )

    # =========================================================================
    # SLIDE 3: Strengthening D365 & Maximizing FTH Value
    # =========================================================================
    s3 = add_slide_scaffolding(
        "Product Excellence & Strategic Value",
        "Strengthening Dynamics 365 & Maximizing Fujitsu (FTH) Value",
        "A strategic technology fusion that drives mutual ROI for enterprise clients and Fujitsu Thailand.",
        3
    )

    col_w = Inches(5.66)
    col_gap = Inches(0.4)

    # Column 1: D365 Product Enhancement
    c_p = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, col_w, card_h)
    c_p.fill.solid()
    c_p.fill.fore_color.rgb = CARD_BG
    c_p.line.color.rgb = CARD_BORDER
    c_p.line.width = Pt(1.2)
    stripe_p = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, col_w, Inches(0.08))
    stripe_p.fill.solid()
    stripe_p.fill.fore_color.rgb = ACCENT_BLUE
    stripe_p.line.fill.background()

    ptf = c_p.text_frame
    ptf.word_wrap = True
    ptf.margin_left = ptf.margin_right = Inches(0.35)
    ptf.margin_top = Inches(0.25)
    p = ptf.paragraphs[0]
    p.text = "💎 1. Dynamics 365 Product Enhancement"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(12)

    p_items = [
        ("Native Copilot Fusion", "Direct integration with Copilot Studio and Azure AI via Microsoft Dataverse—zero external data leakage."),
        ("Autonomous Draft Preparation", "AI extracts data, cross-checks 3-way matching rules, and pre-populates draft vendor invoices—slashing AP review time by >70%."),
        ("Strict Human-in-the-Loop", "AI prepares validated drafts; finance professionals retain exclusive approval authority before any journal posting."),
        ("Total Thai Tax Compliance", "Comprehensive support for PP.30, PND 3/53, 50 Tawi certificates, and Revenue Dept e-Tax XML formats.")
    ]
    for title, desc in p_items:
        pt = ptf.add_paragraph()
        pt.text = f"• {title}: "
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(10.5)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_BODY_HEAD

        run = pt.add_run()
        run.text = desc
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY
        pt.space_after = Pt(7)

    # Column 2: Fujitsu (FTH) Strategic Value
    c_f = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + col_gap, top_pos, col_w, card_h)
    c_f.fill.solid()
    c_f.fill.fore_color.rgb = CARD_BG
    c_f.line.color.rgb = CARD_BORDER
    c_f.line.width = Pt(1.2)
    stripe_f = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + col_gap, top_pos, col_w, Inches(0.08))
    stripe_f.fill.solid()
    stripe_f.fill.fore_color.rgb = ACCENT_RED
    stripe_f.line.fill.background()

    ftf = c_f.text_frame
    ftf.word_wrap = True
    ftf.margin_left = ftf.margin_right = Inches(0.35)
    ftf.margin_top = Inches(0.25)
    p = ftf.paragraphs[0]
    p.text = "🏆 2. Fujitsu (FTH) Enterprise Strategic Value"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p.space_after = Pt(12)

    f_items = [
        ("First-Mover Market Advantage", "FTH becomes the first enterprise partner offering a pre-packaged, AI-infused Thai Localization solution, creating a decisive competitive moat."),
        ("High-Margin Reusable IP Asset", "Packaged as commercial IP, enabling rapid, repeatable deployment across enterprise clients with lower implementation overhead."),
        ("Expands Japanese MNC & Corporate Footprint", "Directly tailored to Japanese multinationals and major Thai conglomerates seeking modern, fully compliant ERP solutions."),
        ("Reinforces Global Trusted Partner Brand", "Combines Fujitsu’s renowned Japanese quality and governance standards with Microsoft’s premier enterprise AI platform.")
    ]
    for title, desc in f_items:
        ft = ftf.add_paragraph()
        ft.text = f"• {title}: "
        ft.font.name = "Segoe UI"
        ft.font.size = Pt(10.5)
        ft.font.bold = True
        ft.font.color.rgb = TEXT_BODY_HEAD

        run = ft.add_run()
        run.text = desc
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY
        ft.space_after = Pt(7)

    # Slide 3 Bottom Banner
    b_bar3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.52))
    b_bar3.fill.solid()
    b_bar3.fill.fore_color.rgb = RGBColor(241, 245, 249)
    b_bar3.line.color.rgb = CARD_BORDER
    b_bar3.line.width = Pt(1)
    btf3 = b_bar3.text_frame
    btf3.word_wrap = True
    btf3.margin_top = Inches(0.12)
    bp3 = btf3.paragraphs[0]
    bp3.text = "🎯 Strategic Impact: Transitioning from bespoke custom implementation to leading the market with proprietary Intellectual Property.  ✓ High Margin  ✓ Scalable  ✓ Repeatable"
    bp3.font.name = "Segoe UI"
    bp3.font.size = Pt(9.5)
    bp3.font.bold = True
    bp3.font.color.rgb = ACCENT_EMERALD
    bp3.alignment = PP_ALIGN.CENTER

    s3.notes_slide.notes_text_frame.text = (
        "【Slide 3: Product Excellence & Fujitsu Strategic Value】\n"
        "Tone: Visionary, High-Impact | Target Duration: ~50 seconds\n\n"
        "\"On Slide 3, we illustrate how this initiative generates substantial, balanced value across both the product platform "
        "and Fujitsu Thailand’s commercial portfolio.\n\n"
        "From a product enhancement perspective, we are embedding Microsoft Copilot Studio and Azure AI directly into Dynamics 365 "
        "via Microsoft Dataverse. This architectural design ensures that enterprise data never leaves your secure tenant. "
        "Rather than replacing human judgment, Copilot functions as an autonomous drafting co-pilot: extracting invoice details, "
        "validating against 3-way matching rules, and pre-populating draft vendor journals. Crucially, we enforce strict Human-in-the-Loop "
        "governance: no transaction is ever committed to the general ledger without explicit review and approval by an authorized accountant.\n\n"
        "From Fujitsu Thailand’s strategic perspective, this initiative establishes a decisive first-mover advantage. "
        "By transforming custom localization scripts into a packaged, repeatable AI-powered Thai Localization IP asset, we significantly "
        "reduce project implementation time while commanding higher consulting margins. This offering directly addresses the acute compliance "
        "needs of Japanese multinationals and major domestic enterprises in Thailand, reinforcing Fujitsu’s legacy of Japanese engineering quality, "
        "robust governance, and trusted delivery.\""
    )

    # =========================================================================
    # SLIDE 4: Executive Agenda & Transition to Live Solution Pitch
    # =========================================================================
    s4 = add_slide_scaffolding(
        "Solution Roadmap & Live Demonstration",
        "Executive Agenda & Transition to Live Solution Pitch",
        "Today’s 12-slide walkthrough and formal handover into the live interactive demonstration platform.",
        4
    )

    agenda_w = Inches(6.8)
    ag_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, agenda_w, card_h)
    ag_card.fill.solid()
    ag_card.fill.fore_color.rgb = CARD_BG
    ag_card.line.color.rgb = CARD_BORDER
    ag_card.line.width = Pt(1.2)
    stripe_ag = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, agenda_w, Inches(0.08))
    stripe_ag.fill.solid()
    stripe_ag.fill.fore_color.rgb = ACCENT_BLUE
    stripe_ag.line.fill.background()

    atf = ag_card.text_frame
    atf.word_wrap = True
    atf.margin_left = atf.margin_right = Inches(0.35)
    atf.margin_top = Inches(0.25)
    p = atf.paragraphs[0]
    p.text = "📑 Key Milestones in Live Interactive Pitch (12 Slides)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_NAVY
    p.space_after = Pt(12)

    agenda_items = [
        ("1. Official Copilot Architecture (Slide 3)", "Multi-Agent Framework and 100% secure enterprise integration with D365 via Dataverse."),
        ("2. Top 5 Thai Localization Use Cases (Slide 4)", "Live walkthrough: AP Invoice Intelligence, Thai Tax Copilot, e-Tax Gateway, Continuous Reconciliation, and Policy Q&A."),
        ("3. 8-Week POC Delivery Plan (Slides 10–11)", "Pragmatic, low-risk deployment roadmap delivering fast measurable value and rapid payback."),
        ("4. Dual-Host AI Narration & Auto-Tour", "Production neural studio voice in English & Thai with automated slide tour capabilities.")
    ]
    for num, txt in agenda_items:
        ap = atf.add_paragraph()
        ap.text = f"• {num}: "
        ap.font.name = "Segoe UI"
        ap.font.size = Pt(10.5)
        ap.font.bold = True
        ap.font.color.rgb = TEXT_BODY_HEAD

        run = ap.add_run()
        run.text = txt
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_BODY
        ap.space_after = Pt(8)

    # Right: Handover Callout Card
    call_w = Inches(4.533)
    call_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + agenda_w + Inches(0.4), top_pos, call_w, card_h)
    call_card.fill.solid()
    call_card.fill.fore_color.rgb = RGBColor(236, 253, 245) # Soft Emerald Tint
    call_card.line.color.rgb = ACCENT_EMERALD
    call_card.line.width = Pt(1.5)

    ctf = call_card.text_frame
    ctf.word_wrap = True
    ctf.margin_left = ctf.margin_right = Inches(0.35)
    ctf.margin_top = Inches(0.3)

    p = ctf.paragraphs[0]
    p.text = "🚀 Ready for Live Demonstration"
    p.font.name = "Segoe UI"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p.space_after = Pt(10)

    p_hand = ctf.add_paragraph()
    p_hand.text = "\"Without further ado, let us now transition directly into the live interactive platform to experience the future of Dynamics 365 and Copilot AI firsthand...\""
    p_hand.font.name = "Segoe UI"
    p_hand.font.size = Pt(10.5)
    p_hand.font.italic = True
    p_hand.font.color.rgb = TEXT_BODY_HEAD
    p_hand.space_after = Pt(14)

    p_btn = ctf.add_paragraph()
    p_btn.text = "🌐 Direct Platform Link (GitHub Pages):"
    p_btn.font.name = "Segoe UI"
    p_btn.font.size = Pt(10)
    p_btn.font.bold = True
    p_btn.font.color.rgb = TEXT_TITLE
    p_btn.space_after = Pt(2)

    p_url = ctf.add_paragraph()
    p_url.text = "https://wititgithub.github.io/fth-d365-fo-copilot-deck/"
    p_url.font.name = "Segoe UI"
    p_url.font.size = Pt(9.5)
    p_url.font.bold = True
    p_url.font.color.rgb = ACCENT_BLUE
    p_url.space_after = Pt(14)

    p_sub_cue = ctf.add_paragraph()
    p_sub_cue.text = "💡 Host Presentation Cue:\nClick the direct URL above or switch to your browser window to start the 12-slide interactive tour."
    p_sub_cue.font.name = "Segoe UI"
    p_sub_cue.font.size = Pt(9)
    p_sub_cue.font.color.rgb = TEXT_MUTED

    # Slide 4 Bottom Banner
    b_bar4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.52))
    b_bar4.fill.solid()
    b_bar4.fill.fore_color.rgb = RGBColor(241, 245, 249)
    b_bar4.line.color.rgb = CARD_BORDER
    b_bar4.line.width = Pt(1)
    btf4 = b_bar4.text_frame
    btf4.word_wrap = True
    btf4.margin_top = Inches(0.12)
    bp4 = btf4.paragraphs[0]
    bp4.text = "🚀 Seamless Handover: Transitioning to the 12-slide Live Interactive Pitch Deck with Production Dual-Host AI Audio."
    bp4.font.name = "Segoe UI"
    bp4.font.size = Pt(9.5)
    bp4.font.bold = True
    bp4.font.color.rgb = ACCENT_BLUE
    bp4.alignment = PP_ALIGN.CENTER

    s4.notes_slide.notes_text_frame.text = (
        "【Slide 4: Executive Agenda & Transition to Live Solution Pitch】\n"
        "Tone: Engaging, Action-Oriented | Target Duration: ~30 seconds\n\n"
        "\"Finally, on Slide 4, let us outline the interactive journey we have prepared for you today.\n\n"
        "Following this opening briefing, we will transition directly into our production-grade Live Interactive Pitch Platform, "
        "which walks through four core pillars across twelve executive slides:\n"
        "1. First, the official Microsoft Copilot Studio reference architecture, detailing enterprise security, Dataverse integration, and multi-agent coordination;\n"
        "2. Second, an in-depth walkthrough of our top five Thai localization use cases—spanning AP Invoice Intelligence, Thai Tax Copilot, e-Tax Gateway pre-validation, Continuous Reconciliation, and Policy Q&A;\n"
        "3. Third, a pragmatic 8-week Proof of Concept delivery roadmap, engineered to deliver measurable ROI and verified business value within two months; and\n"
        "4. Fourth, an interactive presentation experience equipped with dual-host narration and real-time scenario simulation.\n\n"
        "Without further delay, I invite our executive team to transition with me into the live interactive platform to witness the future of Dynamics 365 and Copilot AI in action.\n\n"
        "Please join me as we launch the live demonstration.\""
    )

    out_pptx_paths = [
        r"c:\FTH\FTH-D365_FO\FTH_D365_Executive_Opening_Deck.pptx",
        r"c:\FTH\FTH-D365_FO\assets\FTH_D365_Executive_Opening_Deck.pptx",
        r"C:\Users\witit\.gemini\antigravity\brain\7bc4ceac-b79a-4881-9938-a5db0d7782a2\FTH_D365_Executive_Opening_Deck.pptx"
    ]

    for p in out_pptx_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        prs.save(p)
        print(f"Generated Executive PowerPoint Deck -> {p} ({os.path.getsize(p)//1024} KB)")

if __name__ == "__main__":
    build_presentation()
